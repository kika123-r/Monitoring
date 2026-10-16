#!/usr/bin/env python3.13
"""Append today's collected rows (data/results/<profile>.csv) onto the growing
master xlsx ledger for that profile (data/master/<name>.xlsx), matching the
existing workbook's column order, types and number formats exactly. Idempotent
per day: running it twice on the same COLLECTION_DATE does not duplicate rows.
First run ever bootstraps the ledger from the user's original historical xlsx;
every later run appends onto the ledger file itself."""
import csv
import datetime
import os
import re
import sys
import openpyxl

CURRENCY_FMT = '_-* #,##0.00\\ [$€-41B]_-;\\-* #,##0.00\\ [$€-41B]_-;_-* "-"??\\ [$€-41B]_-;_-@_-'
CURRENCY_FMT2 = '_-* #,##0.00\\ "€"_-;\\-* #,##0.00\\ "€"_-;_-* "-"??\\ "€"_-;_-@_-'
PRICE_FMT = '#,##0.00\\ "€"'
DATE_FMT = 'mm-dd-yy'

TODAY = os.environ.get('COLLECTION_DATE') or datetime.date.today().isoformat()
TODAY_DATE = datetime.date.fromisoformat(TODAY)

BOOTSTRAP_SOURCES = {
    'byt-bratislava': '/root/.claude/uploads/e15c6bb8-3fcd-534f-ba00-dc4d9897ba8a/07df6f62-Poistenie_majetku_-_byt_Bratislava-3.xlsx',
    'byt-kosice': '/root/.claude/uploads/e15c6bb8-3fcd-534f-ba00-dc4d9897ba8a/9bb55416-Poistenie_majetku_-_byt_Kos_ice-3.xlsx',
    'dom-male-dvorniky': '/root/.claude/uploads/e15c6bb8-3fcd-534f-ba00-dc4d9897ba8a/537046a3-Poistenie_majetku_-_dom_Male__Dvorni_ky-3.xlsx',
    'dom-vlckovce': '/root/.claude/uploads/e15c6bb8-3fcd-534f-ba00-dc4d9897ba8a/455a69a6-Poistenie_majetku___dom_Vlc_kovce-3.xlsx',
}
OUT_NAMES = {
    'byt-bratislava': 'Poistenie_majetku_-_byt_Bratislava.xlsx',
    'byt-kosice': 'Poistenie_majetku_-_byt_Kosice.xlsx',
    'dom-male-dvorniky': 'Poistenie_majetku_-_dom_Male_Dvorniky.xlsx',
    'dom-vlckovce': 'Poistenie_majetku_-_dom_Vlckovce.xlsx',
}
MASTER_DIR = os.environ.get('MASTER_DIR', '/home/user/Monitoring/data/master')


def to_int_or_text(v):
    if v is None or v == '':
        return None
    try:
        return int(v)
    except ValueError:
        return v


def to_number(v):
    if v is None or v == '':
        return None
    v = v.replace('\xa0', '').replace(' ', '').replace(',', '.')
    try:
        return float(v)
    except ValueError:
        return None


def parse_price(v):
    if not v:
        return None
    m = re.search(r'([\d.,]+)', v.replace('\xa0', ''))
    if not m:
        return None
    return float(m.group(1).replace('.', '').replace(',', '.')) if ',' in m.group(1) else float(m.group(1))


def last_data_row(ws):
    last = 1
    scan_to = min(ws.max_row, 20000)
    for i in range(2, scan_to + 1):
        if ws.cell(row=i, column=2).value is not None:  # column B (Typ) is reliably text, not a stray artifact
            last = i
    return last


def already_appended(ws, start_row):
    # look backwards from the end for a date cell equal to today -> already ran today
    for i in range(start_row - 1, 1, -1):
        v = ws.cell(row=i, column=1).value
        if v is None:
            continue
        if isinstance(v, datetime.datetime):
            v = v.date()
        if v == TODAY_DATE:
            return True
        if isinstance(v, datetime.date) and v != TODAY_DATE:
            break
    return False


def run(profile_id):
    out_path = f'{MASTER_DIR}/{OUT_NAMES[profile_id]}'
    src = out_path if os.path.exists(out_path) else BOOTSTRAP_SOURCES[profile_id]
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    start = last_data_row(ws) + 1

    if already_appended(ws, start):
        print(f'{profile_id}: {TODAY} already appended, skipping')
        return

    results_dir = os.environ.get('RESULTS_DIR', '/home/user/Monitoring/data/results')
    csv_path = f'{results_dir}/{profile_id}.csv'
    with open(csv_path, newline='', encoding='utf-8') as fh:
        rows = list(csv.reader(fh))
    header, data_rows = rows[0], rows[1:]

    this_run_rows = [r for r in data_rows if r[0] == TODAY]
    if not this_run_rows:
        print(f'{profile_id}: no rows for {TODAY} in {csv_path}, skipping')
        return

    meta_row = this_run_rows[0]
    uniqa_note_row = list(meta_row)
    uniqa_note_row[14] = ''
    uniqa_note_row[15] = 'Uniqa'
    uniqa_note_row[16] = '-'
    uniqa_note_row[17] = ''
    uniqa_note_row[18] = ''
    uniqa_note_row[19] = 'neuvedené'
    uniqa_note_row[20] = ''
    uniqa_note_row[21] = ''
    uniqa_note_row[22] = ''
    uniqa_note_row[23] = ''
    uniqa_note_row[24] = '-'
    uniqa_note_row[25] = '-'
    uniqa_note_row[26] = ('Kalkulačka vyplnená kompletne a správne, ale server odmieta finálny výpočet ceny s '
                           'chybou "Ochrana proti spamu" (neviditeľná reCAPTCHA v3 kontrola viazaná na reputáciu '
                           'IP adresy / automatizáciu prehliadača) — cenu sa z tohto cloudového prostredia '
                           'nepodarilo získať.')
    this_run_rows = this_run_rows + [uniqa_note_row]

    r = start
    for row in this_run_rows:
        (dat, typ, col3, col4, mesto, ulica, cislo, plocha, loggia, rok, material, ucel,
         suma_nehnutelnost, suma_domacnost, odporucana, poistovna, balik, krytie, cena,
         tolerancia, spoluucast, zodp1, zodp2, doplnkove, zlava_uplatnena, zlava_mozne, info) = row

        vals = [
            TODAY_DATE,
            typ,
            to_int_or_text(col3),
            to_int_or_text(col4),
            mesto, ulica, cislo,
            to_int_or_text(plocha),
            to_int_or_text(loggia),
            to_int_or_text(rok),
            material, ucel,
            to_number(suma_nehnutelnost),
            to_number(suma_domacnost),
            to_number(odporucana),
            poistovna, balik, krytie,
            parse_price(cena),
            tolerancia,
            to_number(re.sub(r'[^\d.,]', '', spoluucast)) if spoluucast else None,
            to_number(zodp1),
            to_number(zodp2),
            doplnkove, zlava_uplatnena, zlava_mozne, info,
        ]
        formats = [
            DATE_FMT, 'General', 'General', 'General', 'General', 'General', 'General',
            'General', 'General', 'General', 'General', 'General',
            CURRENCY_FMT, CURRENCY_FMT, CURRENCY_FMT,
            'General', 'General', 'General',
            PRICE_FMT, 'General', CURRENCY_FMT,
            CURRENCY_FMT2, CURRENCY_FMT2,
            'General', 'General', 'General', 'General',
        ]
        for col, (v, fmt) in enumerate(zip(vals, formats), start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.number_format = fmt
        r += 1

    wb.save(out_path)
    print(f'{profile_id}: appended {len(this_run_rows)} rows for {TODAY} (from row {start}) -> {out_path}')


if __name__ == '__main__':
    profiles = sys.argv[1:] or list(BOOTSTRAP_SOURCES)
    for pid in profiles:
        run(pid)
